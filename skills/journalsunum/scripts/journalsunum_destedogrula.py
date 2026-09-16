#!/usr/bin/env python3
# Adapted from k-dense-ai/scientific-agent-skills/skills/scientific-slides/scripts/validate_presentation.py, MIT, K-Dense Inc. Renamed to this package's N12 rule; imports re-pointed; 1.20.0 added the text-budget pass (journalsunum_metinolcer), --json/--output/--thresholds and exit 2 for an unreadable package; 1.21.0 made file size a warning and added --privacy (identifier review scan) (nothing else changed).
"""
Presentation Validation Script

Validates scientific presentations for common issues:
- Slide count vs. duration
- Slide text budget against journalsunum-r-tasarim.md §2 (bullets, words, fonts, notes)
- LaTeX compilation
- File size checks
- Basic format validation

Exit codes: 0 pass · 1 issues found (or file missing) · 2 the package cannot be read.
"""

import sys
import os
import re
import argparse
import subprocess
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Dict, List, Tuple, Optional

# PDF page and geometry analysis. pypdf is the maintained continuation of
# PyPDF2 and exposes the same PdfReader API, so accept whichever is installed.
try:
    from pypdf import PdfReader
    HAS_PDF_READER = True
except ImportError:
    try:
        from PyPDF2 import PdfReader
        HAS_PDF_READER = True
    except ImportError:
        HAS_PDF_READER = False

# Try to import python-pptx for PowerPoint analysis
try:
    from pptx import Presentation
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False


class PresentationValidator:
    """Validates presentations for common issues."""
    
    # Recommended slide counts by duration (min, recommended, max)
    SLIDE_GUIDELINES = {
        5: (5, 6, 8),
        10: (8, 11, 14),
        15: (13, 16, 20),
        20: (18, 22, 26),
        30: (22, 27, 33),
        45: (32, 40, 50),
        60: (40, 52, 65),
    }
    
    # Suffixes the dependency-free text-budget pass can read (OOXML packages).
    TEXT_BUDGET_SUFFIXES = ('.pptx', '.potx', '.ppsx')

    def __init__(
        self,
        filepath: str,
        duration: Optional[int] = None,
        text_budget: bool = True,
        thresholds: Optional[Dict] = None,
        privacy: bool = False,
    ):
        self.filepath = Path(filepath)
        self.duration = duration
        self.text_budget = text_budget
        self.thresholds = thresholds
        self.privacy = privacy
        self.privacy_report: Optional[Dict] = None
        self.text_budget_report: Optional[Dict] = None
        self.file_type = self.filepath.suffix.lower()
        self.issues = []
        self.warnings = []
        self.info = []

    def validate(self, banner: bool = True) -> Dict:
        """Run all validations and return results."""
        if banner:
            print(f"Validating: {self.filepath.name}")
            print(f"File type: {self.file_type}")
            print("=" * 60)
        
        # Check file exists
        if not self.filepath.exists():
            self.issues.append(f"File not found: {self.filepath}")
            return self._format_results()
        
        # File size check
        self._check_file_size()
        
        # Type-specific validation
        if self.file_type == '.pdf':
            self._validate_pdf()
        elif self.file_type in ['.pptx', '.ppt', '.potx', '.ppsx']:
            self._validate_pptx()
        elif self.file_type in ['.tex']:
            self._validate_latex()
        else:
            self.warnings.append(f"Unknown file type: {self.file_type}")
        
        return self._format_results()
    
    def _check_file_size(self):
        """Check if file size is reasonable."""
        size_mb = self.filepath.stat().st_size / (1024 * 1024)
        self.info.append(f"File size: {size_mb:.2f} MB")
        
        # Size is advice, never a gate: exit 1 is reserved for a broken ceiling, and a
        # deck with embedded clinical video is always over 100 MB (observation 163).
        if size_mb > 100:
            self.warnings.append(
                f"File is very large ({size_mb:.1f} MB). "
                "Consider compressing images or trimming embedded video."
            )
        elif size_mb > 50:
            self.warnings.append(
                f"File is large ({size_mb:.1f} MB). "
                "May be slow to email or upload."
            )
    
    def _validate_pdf(self):
        """Validate PDF presentation."""
        if not HAS_PDF_READER:
            self.warnings.append(
                "pypdf not installed. Install with: uv pip install pypdf"
            )
            return

        try:
            with open(self.filepath, 'rb') as f:
                reader = PdfReader(f)
                num_pages = len(reader.pages)
                
                self.info.append(f"Number of slides: {num_pages}")
                
                # Check slide count against duration
                if self.duration:
                    self._check_slide_count(num_pages)
                
                # Get page size
                first_page = reader.pages[0]
                media_box = first_page.mediabox
                width = float(media_box.width)
                height = float(media_box.height)
                
                # Convert points to inches (72 points = 1 inch)
                width_in = width / 72
                height_in = height / 72
                aspect = width / height
                
                self.info.append(
                    f"Slide dimensions: {width_in:.1f}\" × {height_in:.1f}\" "
                    f"(aspect ratio: {aspect:.2f})"
                )
                
                # Check common aspect ratios
                if abs(aspect - 16/9) < 0.01:
                    self.info.append("Aspect ratio: 16:9 (widescreen)")
                elif abs(aspect - 4/3) < 0.01:
                    self.info.append("Aspect ratio: 4:3 (standard)")
                else:
                    self.warnings.append(
                        f"Unusual aspect ratio: {aspect:.2f}. "
                        "Confirm this matches venue requirements."
                    )
                
        except Exception as e:
            self.issues.append(f"Error reading PDF: {str(e)}")
    
    def _validate_pptx(self):
        """Validate PowerPoint presentation.

        The text-budget pass (1.20.0) reads the package as ZIP/XML and needs no pip
        package, so it also supplies the slide count; python-pptx, when present, only
        adds the slide dimensions. A CliError from the reader means the package itself
        is unreadable — it propagates to main() as exit 2 rather than becoming a finding.
        """
        num_slides = None

        if self.text_budget and self.file_type in self.TEXT_BUDGET_SUFFIXES:
            from journalsunum_metinolcer import measure_deck, summarize

            report = measure_deck(self.filepath, thresholds=self.thresholds)
            self.text_budget_report = report
            num_slides = report['totals']['slides']
            self.info.append(f"Number of slides: {num_slides}")
            info, warnings, issues = summarize(report)
            self.info.extend(info)
            self.warnings.extend(warnings)
            self.issues.extend(issues)

        if not HAS_PPTX:
            if num_slides is None:
                self.warnings.append(
                    "python-pptx not installed. Install with: uv pip install python-pptx"
                )
                return
        else:
            try:
                prs = Presentation(self.filepath)
                if num_slides is None:
                    num_slides = len(prs.slides)
                    self.info.append(f"Number of slides: {num_slides}")

                # Get slide dimensions
                width_inches = prs.slide_width / 914400  # EMU to inches
                height_inches = prs.slide_height / 914400
                aspect = prs.slide_width / prs.slide_height

                self.info.append(
                    f"Slide dimensions: {width_inches:.1f}\" × {height_inches:.1f}\" "
                    f"(aspect ratio: {aspect:.2f})"
                )

                # Legacy font/bullet heuristics, superseded by the text-budget pass.
                if self.text_budget_report is None:
                    self._check_pptx_content(prs)

            except Exception as e:
                self.issues.append(f"Error reading PowerPoint: {str(e)}")

        # Check slide count against duration
        if self.duration and num_slides is not None:
            self._check_slide_count(num_slides)

        if self.privacy and self.file_type in self.TEXT_BUDGET_SUFFIXES:
            self._check_privacy()

    # --privacy (1.21.0, observation 162). A clinical draft carries identifiers that no
    # rendered view shows: DICOM overlay text burned into 16-bit images (PowerPoint draws
    # them clipped), patient names left in shape names and alt text, initials and national
    # ID numbers in body text or notes. The scan reads bytes and package metadata, never
    # the render. Every finding is a REVIEW item for a person: it never adds an issue, so
    # the exit code keeps meaning "a text-budget ceiling is broken".
    PRIVACY_LETTERS = 'A-Za-zÇĞİÖŞÜçğıöşü'
    PRIVACY_INITIALS = re.compile(
        r'(?<![A-Za-zÇĞİÖŞÜçğıöşü])[A-ZÇĞİÖŞÜ]\.\s?[A-ZÇĞİÖŞÜ]\.?(?![A-Za-zÇĞİÖŞÜçğıöşü])'
    )
    PRIVACY_NATIONAL_ID = re.compile(r'(?<!\d)[1-9]\d{10}(?!\d)')
    PRIVACY_RECORD_WORDS = re.compile(
        r'\b(PAT\s?\d{3,}|MRN\s?:?\s?\d+|protokol\s*(no|numaras[ıi])|hasta\s*(no|numaras[ıi]|ad[ıi])|dosya\s*no)\b',
        re.IGNORECASE,
    )
    PRIVACY_GENERIC_NAME = re.compile(
        r'^(Picture|Resim|Image|Görüntü|Title|Başlık|Unvan|Subtitle|Alt Başlık|Text|Metin|'
        r'sketch line|Content Placeholder|'
        r'İçerik Yer Tutucusu|Text Placeholder|Metin Yer Tutucusu|TextBox|Text Box|Metin Kutusu|'
        r'Slide Number Placeholder|Slayt Numarası Yer Tutucusu|Footer Placeholder|'
        r'Alt Bilgi Yer Tutucusu|Date Placeholder|Tarih Yer Tutucusu|Rectangle|Dikdörtgen|Oval|'
        r'Straight Connector|Düz Bağlayıcı|Straight Arrow Connector|Group|Grup|Table|Tablo|Chart|'
        r'Grafik|Media|Ortam|Video|Online Media|Shape|Şekil|Line|Çizgi|Arrow|Ok|Freeform|Serbest Form)'
        r'[\s_\-]*\d*$',
        re.IGNORECASE,
    )
    PRIVACY_IMAGE_SUFFIXES = ('.tif', '.tiff', '.dcm', '.png', '.jpg', '.jpeg', '.bmp')
    PRIVACY_XML_LIMIT = 20 * 1024 * 1024
    NS = {
        'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
        'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
        'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
        'dc': 'http://purl.org/dc/elements/1.1/',
        'cp': 'http://schemas.openxmlformats.org/package/2006/metadata/core-properties',
    }

    def _privacy_slide_order(self, zf) -> Dict[str, int]:
        """Map ppt/slides/slideN.xml to its position in the presentation (not its file number)."""
        order: Dict[str, int] = {}
        try:
            rels = ET.fromstring(zf.read('ppt/_rels/presentation.xml.rels'))
            targets = {
                rel.get('Id'): 'ppt/' + rel.get('Target', '').lstrip('/').replace('ppt/', '', 1)
                for rel in rels
            }
            pres = ET.fromstring(zf.read('ppt/presentation.xml'))
            ids = pres.find('p:sldIdLst', self.NS)
            for pos, sld in enumerate(ids if ids is not None else [], start=1):
                target = targets.get(sld.get('{%s}id' % self.NS['r']))
                if target:
                    order[target] = pos
        except (KeyError, ET.ParseError):
            pass
        return order

    def _check_privacy(self):
        findings: List[Dict] = []

        def add(code, message, value=None, slide=None, part=None):
            item = {'code': code, 'severity': 'review', 'message': message}
            if slide is not None:
                item['slide'] = slide
            if part:
                item['part'] = part
            if value is not None:
                item['value'] = str(value)[:80]
            findings.append(item)

        def scan_text(text, slide, part):
            for m in self.PRIVACY_INITIALS.finditer(text):
                add('PRIVACY_INITIALS', 'initials pattern in text — patient initials?', m.group(0), slide, part)
            for m in self.PRIVACY_NATIONAL_ID.finditer(text):
                add('PRIVACY_NATIONAL_ID', '11-digit number in text — national ID number?', m.group(0), slide, part)
            for m in self.PRIVACY_RECORD_WORDS.finditer(text):
                add('PRIVACY_RECORD_REFERENCE', 'record/patient reference in text', m.group(0), slide, part)

        try:
            zf = zipfile.ZipFile(self.filepath)
        except zipfile.BadZipFile:
            add('PRIVACY_UNREADABLE', 'package is not a ZIP; privacy scan skipped')
            self.privacy_report = {'findings': findings}
            return

        with zf:
            order = self._privacy_slide_order(zf)
            names = zf.namelist()
            slide_parts = [n for n in names if re.match(r'ppt/slides/slide\d+\.xml$', n)]
            for part in sorted(slide_parts, key=lambda n: order.get(n, 10 ** 6)):
                slide = order.get(part)
                info = zf.getinfo(part)
                if info.file_size > self.PRIVACY_XML_LIMIT:
                    add('PRIVACY_PART_TOO_LARGE', 'slide XML too large to scan', slide=slide, part=part)
                    continue
                try:
                    root = ET.fromstring(zf.read(part))
                except ET.ParseError:
                    add('PRIVACY_UNREADABLE', 'slide XML does not parse', slide=slide, part=part)
                    continue
                texts = [t.text or '' for t in root.iter('{%s}t' % self.NS['a'])]
                scan_text(' '.join(texts), slide, part)
                for el in root.iter():
                    if not el.tag.endswith('}cNvPr'):
                        continue
                    name = (el.get('name') or '').strip()
                    descr = (el.get('descr') or '').strip()
                    title = (el.get('title') or '').strip()
                    if name and not self.PRIVACY_GENERIC_NAME.match(name):
                        add('PRIVACY_SHAPE_NAME', 'non-generic shape name — may carry a person or file name', name, slide, part)
                    for value in (descr, title):
                        if value:
                            add('PRIVACY_ALT_TEXT', 'alt text / title on a shape — check for identifiers', value, slide, part)

            notes_parts = [n for n in names if re.match(r'ppt/notesSlides/notesSlide\d+\.xml$', n)]
            for part in notes_parts:
                try:
                    root = ET.fromstring(zf.read(part))
                except (ET.ParseError, KeyError):
                    continue
                scan_text(' '.join(t.text or '' for t in root.iter('{%s}t' % self.NS['a'])), None, part)

            if 'docProps/core.xml' in names:
                try:
                    core = ET.fromstring(zf.read('docProps/core.xml'))
                    for tag in ('dc:creator', 'cp:lastModifiedBy', 'dc:title', 'dc:subject'):
                        el = core.find(tag, self.NS)
                        if el is not None and (el.text or '').strip():
                            add('PRIVACY_PACKAGE_METADATA', f'package metadata {tag}', el.text.strip(), part='docProps/core.xml')
                except ET.ParseError:
                    pass

            try:
                from PIL import Image
                has_pil = True
            except ImportError:
                has_pil = False
            for part in names:
                lower = part.lower()
                if lower.startswith('ppt/media/') and lower.endswith(('.mp4', '.mov', '.avi', '.wmv', '.m4v')):
                    add('PRIVACY_VIDEO_NOT_INSPECTED',
                        'embedded video — only its poster frame is readable here; frames beyond it are NOT inspected',
                        part=part)
            for part in names:
                lower = part.lower()
                if not (lower.startswith('ppt/media/') and lower.endswith(self.PRIVACY_IMAGE_SUFFIXES)):
                    continue
                if lower.endswith('.dcm'):
                    add('PRIVACY_DICOM_FILE', 'embedded DICOM file — carries patient tags by design', part=part)
                    continue
                if lower.endswith(('.tif', '.tiff')):
                    add('PRIVACY_IMAGE_OVERLAY_RISK',
                        'TIFF image — typical export of fluoroscopy/DICOM; burned-in overlay text may be present, hidden or fully visible depending on the renderer',
                        part=part)
                if not has_pil:
                    continue
                try:
                    import io
                    with Image.open(io.BytesIO(zf.read(part))) as im:
                        if im.mode in ('I;16', 'I;16B', 'I;16L', 'I', 'F'):
                            add('PRIVACY_IMAGE_HIGH_BIT_DEPTH',
                                f'{im.mode} image — an overlay may be hidden OR fully visible depending on the renderer; '
                                'inspect the rendered slide and the normalised pixels',
                                part=part)
                        meta = getattr(im, 'tag_v2', None)
                        if meta:
                            for tag_id, label in ((270, 'ImageDescription'), (315, 'Artist'), (305, 'Software')):
                                value = meta.get(tag_id)
                                if value and tag_id != 305:
                                    add('PRIVACY_IMAGE_METADATA', f'TIFF tag {label}', value, part=part)
                        exif = im.getexif() if hasattr(im, 'getexif') else None
                        if exif:
                            for tag_id, label in ((270, 'ImageDescription'), (315, 'Artist'), (37510, 'UserComment')):
                                value = exif.get(tag_id)
                                if value:
                                    add('PRIVACY_IMAGE_METADATA', f'EXIF {label}', value, part=part)
                except Exception:
                    continue

        # One finding per (code, value, part) — a repeated shape name is reported once.
        seen = set()
        unique = []
        for item in findings:
            key = (item['code'], item.get('value'), item.get('part'))
            if key not in seen:
                seen.add(key)
                unique.append(item)
        counts: Dict[str, int] = {}
        for item in unique:
            counts[item['code']] = counts.get(item['code'], 0) + 1
        self.privacy_report = {
            'basis': 'review items for a person; nothing here is a verdict and nothing changes the exit code',
            'counts': counts,
            'findings': unique,
        }
        if unique:
            summary = ', '.join(f'{code} {n}' for code, n in sorted(counts.items()))
            self.warnings.append(f"Privacy: {len(unique)} item(s) to review before reuse ({summary})")
        else:
            self.info.append("Privacy: no identifier patterns found (a clean scan is not proof of de-identification)")

    def _check_pptx_content(self, prs):
        """Check PowerPoint content for common issues."""
        small_text_slides = []
        many_bullets_slides = []
        
        for idx, slide in enumerate(prs.slides, start=1):
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                
                text_frame = shape.text_frame
                
                # Check for small fonts
                for paragraph in text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.font.size and run.font.size.pt < 18:
                            small_text_slides.append(idx)
                            break
                
                # Check for too many bullets
                bullet_count = sum(1 for p in text_frame.paragraphs if p.level == 0)
                if bullet_count > 6:
                    many_bullets_slides.append(idx)
        
        # Report issues
        if small_text_slides:
            unique_slides = sorted(set(small_text_slides))
            self.warnings.append(
                f"Small text (<18pt) found on slides: {unique_slides[:5]}"
                + (" ..." if len(unique_slides) > 5 else "")
            )
        
        if many_bullets_slides:
            unique_slides = sorted(set(many_bullets_slides))
            self.warnings.append(
                f"Many bullets (>6) on slides: {unique_slides[:5]}"
                + (" ..." if len(unique_slides) > 5 else "")
            )
    
    def _validate_latex(self):
        """Validate LaTeX Beamer presentation."""
        self.info.append("LaTeX source file detected")
        
        # Try to compile
        if self._try_compile_latex():
            self.info.append("LaTeX compilation: SUCCESS")
            
            # If PDF was generated, validate it
            pdf_path = self.filepath.with_suffix('.pdf')
            if pdf_path.exists():
                pdf_validator = PresentationValidator(str(pdf_path), self.duration)
                pdf_results = pdf_validator.validate()
                
                # Merge results
                self.info.extend(pdf_results['info'])
                self.warnings.extend(pdf_results['warnings'])
                self.issues.extend(pdf_results['issues'])
        else:
            self.issues.append(
                "LaTeX compilation failed. Check .log file for errors."
            )
    
    def _try_compile_latex(self) -> bool:
        """Try to compile LaTeX file."""
        try:
            # Try pdflatex
            result = subprocess.run(
                ['pdflatex', '-no-shell-escape', '-interaction=nonstopmode', self.filepath.name],
                cwd=self.filepath.parent,
                capture_output=True,
                timeout=60
            )
            return result.returncode == 0
        except (subprocess.TimeoutExpired, FileNotFoundError):
            return False
    
    def _check_slide_count(self, num_slides: int):
        """Check if slide count is appropriate for duration."""
        if self.duration not in self.SLIDE_GUIDELINES:
            # Find nearest duration
            durations = sorted(self.SLIDE_GUIDELINES.keys())
            nearest = min(durations, key=lambda x: abs(x - self.duration))
            min_slides, rec_slides, max_slides = self.SLIDE_GUIDELINES[nearest]
            self.info.append(
                f"Using guidelines for {nearest}-minute talk "
                f"(closest to {self.duration} minutes)"
            )
        else:
            min_slides, rec_slides, max_slides = self.SLIDE_GUIDELINES[self.duration]
        
        self.info.append(
            f"Recommended slides for {self.duration}-minute talk: "
            f"{min_slides}-{max_slides} (optimal: ~{rec_slides})"
        )
        
        if num_slides < min_slides:
            self.warnings.append(
                f"Fewer slides ({num_slides}) than recommended ({min_slides}-{max_slides}). "
                "May have too much time or too little content."
            )
        elif num_slides > max_slides:
            self.warnings.append(
                f"More slides ({num_slides}) than recommended ({min_slides}-{max_slides}). "
                "Likely to run over time."
            )
        else:
            self.info.append(
                f"Slide count ({num_slides}) is within recommended range."
            )
    
    def _format_results(self) -> Dict:
        """Format validation results."""
        return {
            'filepath': str(self.filepath),
            'file_type': self.file_type,
            'info': self.info,
            'warnings': self.warnings,
            'issues': self.issues,
            'valid': len(self.issues) == 0,
            'text_budget': self.text_budget_report,
            'privacy': self.privacy_report,
        }


def _utf8_stdout():
    """The prose report prints emoji; a Windows console codepage would abort on them."""
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding='utf-8', errors='replace')
        except (AttributeError, ValueError):
            pass


def print_results(results: Dict):
    """Print validation results in a readable format."""
    print()
    print("=" * 60)
    print("VALIDATION RESULTS")
    print("=" * 60)
    
    # Print info
    if results['info']:
        print("\n📋 Information:")
        for item in results['info']:
            print(f"  • {item}")
    
    # Print warnings
    if results['warnings']:
        print("\n⚠️  Warnings:")
        for item in results['warnings']:
            print(f"  • {item}")
    
    # Print issues
    if results['issues']:
        print("\n❌ Issues:")
        for item in results['issues']:
            print(f"  • {item}")
    
    # Overall status
    print("\n" + "=" * 60)
    if results['valid']:
        print("✅ Validation PASSED")
        if results['warnings']:
            print(f"   ({len(results['warnings'])} warning(s) found)")
    else:
        print("❌ Validation FAILED")
        print(f"   ({len(results['issues'])} issue(s) found)")
    print("=" * 60)


def main():
    _utf8_stdout()
    parser = argparse.ArgumentParser(
        description='Validate scientific presentations',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  %(prog)s presentation.pdf --duration 15
  %(prog)s slides.pptx --duration 45
  %(prog)s beamer_talk.tex --duration 20

Supported file types:
  - PDF (.pdf)
  - PowerPoint (.pptx, .ppt)
  - LaTeX Beamer (.tex)

Validation checks:
  - Slide count vs. duration
  - File size
  - Slide dimensions
  - Font sizes (PowerPoint)
  - LaTeX compilation (Beamer)
        """
    )
    
    parser.add_argument(
        'filepath',
        help='Path to presentation file (PDF, PPTX, or TEX)'
    )
    
    parser.add_argument(
        '--duration', '-d',
        type=int,
        help='Presentation duration in minutes'
    )
    
    parser.add_argument(
        '--quiet', '-q',
        action='store_true',
        help='Only show issues and warnings'
    )

    parser.add_argument(
        '--text-budget',
        dest='text_budget',
        action='store_true',
        default=True,
        help='Measure slide text against journalsunum-r-tasarim.md §2 (default for .pptx)'
    )

    parser.add_argument(
        '--no-text-budget',
        dest='text_budget',
        action='store_false',
        help='Skip the text-budget pass (a one-slide poster has no bullet budget)'
    )

    parser.add_argument(
        '--thresholds',
        help='JSON file overriding the §2 defaults (same keys as the report prints)'
    )

    parser.add_argument(
        '--privacy',
        action='store_true',
        help='Scan the package for patient identifiers (initials, ID numbers, shape names, '
             'alt text, metadata, high-bit-depth/TIFF images); review items, never an exit-code failure'
    )

    parser.add_argument(
        '--json',
        dest='as_json',
        action='store_true',
        help='Emit the full report as JSON instead of the prose summary'
    )

    parser.add_argument(
        '--output',
        help='Write the JSON report to this new path (implies --json)'
    )

    args = parser.parse_args()

    from journalsunum_ortak import CliError, emit_json, load_json_file

    try:
        thresholds = None
        if args.thresholds:
            _, thresholds = load_json_file(args.thresholds)

        # Validate
        validator = PresentationValidator(
            args.filepath, args.duration,
            text_budget=args.text_budget, thresholds=thresholds,
            privacy=args.privacy,
        )
        results = validator.validate(banner=not (args.as_json or args.output))
    except CliError as exc:
        parser.exit(2, f"error: {exc}\n")

    # Print results
    if args.as_json or args.output:
        emit_json(results, output=args.output)
    elif args.quiet:
        # Only show warnings and issues
        if results['warnings'] or results['issues']:
            print_results(results)
        else:
            print("✅ No issues found")
    else:
        print_results(results)

    # Exit with appropriate code
    sys.exit(0 if results['valid'] else 1)


if __name__ == '__main__':
    main()

